import io

from docx import Document
from openpyxl import load_workbook

from nexus.deliverables import render
from tests.conftest import signup

MEMO = """# Memorandum

**To:** Priya Chakravarti, Partner

## 1. Change of control

The Globex contract requires consent on a change of control [S1]. Consent has not been obtained [S2].

- Risk: termination right on closing [S1]
1. Seek consent before signing

| Contract | Consent needed |
|---|---|
| Globex | Yes [S1] |
"""

SOURCES = [
    {"id": "S1", "verified": True, "quote": "consent of the other party", "document": "Globex MSA.docx"},
    {"id": "S2", "verified": False, "quote": "made up", "reason": "not found"},
]


def test_memo_renders_to_word_with_numbered_sources():
    data, content_type = render("memo.docx", MEMO, SOURCES)
    assert content_type.endswith("wordprocessingml.document")
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Change of control" in text and "[1]" in text
    assert "[unverified]" in text  # S2 failed verification and is shown as such
    assert "Globex MSA.docx" in text and "Sources" in text
    assert doc.tables[0].cell(1, 0).text == "Globex"
    to_line = next(p for p in doc.paragraphs if p.text.startswith("To:"))
    assert to_line.runs[0].bold and to_line.runs[0].text == "To:"


def test_tables_render_to_excel_sheets():
    data, _ = render("schedule.xlsx", "# Consents\n| Contract | Status |\n|---|---|\n| Globex | Pending [S1] |", SOURCES)
    sheet = load_workbook(io.BytesIO(data))["Consents"]
    assert [c.value for c in sheet[2]] == ["Globex", "Pending"]


def test_draft_endpoint_validates_file_names(client):
    signup(client)
    matter = client.post("/api/matters", json={"name": "m"}).json()
    r = client.post(f"/api/matters/{matter['id']}/drafts",
                    json={"instructions": "Draft a memo on consents.", "deliverables": ["../../etc/passwd"]})
    assert r.status_code in (409, 422)  # no documents yet, or bad name
