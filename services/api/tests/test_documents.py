import io

from docx import Document

from nexus.ingest.chunk import chunk
from nexus.ingest.extract import DOCX, TEXT, Page, extract
from nexus.redline import locate, tracked_changes
from nexus.samples import sample_nda_docx, sample_nda_text


def test_docx_extracts_every_clause():
    pages = extract(sample_nda_docx(), DOCX)
    text = pages[0].text
    assert "Mutual Non-Disclosure Agreement" in text
    assert "twenty-four (24) months" in text


def test_passages_carry_their_heading():
    passages = chunk(extract(sample_nda_docx(), DOCX))
    non_solicit = next(p for p in passages if "solicit, employ or engage" in p.text)
    assert non_solicit.heading.startswith("6. Non-Solicitation")
    assert all(len(p.text) <= 2000 for p in passages)
    assert [p.seq for p in passages] == list(range(len(passages)))


def test_text_extract_and_long_paragraphs_split():
    long_paragraph = " ".join(["The Supplier shall deliver the Goods on time."] * 120)
    passages = chunk([Page(1, sample_nda_text() + "\n\n" + long_paragraph)])
    assert max(len(p.text) for p in passages) <= 2000
    assert extract(b"hello", TEXT)[0].text == "hello"


def test_locate_ignores_typography():
    text = "The obligations in clause 3 survive for one (1) year after expiry."
    assert locate("survive for one (1) year", text) == (text.index("survive"), text.index(" after"))
    assert locate("clause 3 … after expiry", text) is not None
    assert locate("two (2) years", text) is None


def test_tracked_changes_are_real_word_revisions():
    findings = [
        {"title": "Duration of the obligations", "status": "deviates",
         "explanation": "Obligations end one year after termination.",
         "suggested_language": "survive for three (3) years after expiry or termination",
         "citation": {"quote": "survive for one (1) year after expiry or termination"}},
        {"title": "Disclosure required by law", "status": "missing",
         "explanation": "No compelled-disclosure permission.",
         "suggested_language": "The Receiving Party may disclose Confidential Information to the extent required by law.",
         "citation": None},
        {"title": "Definition", "status": "meets", "explanation": "", "suggested_language": "", "citation": None},
    ]
    data, placed = tracked_changes(sample_nda_docx(), findings)
    assert placed == {"applied": 2, "not_placed": []}
    doc = Document(io.BytesIO(data))
    xml = doc.element.xml
    assert xml.count("<w:del ") == 1 and xml.count("<w:ins ") == 2
    assert "Nexus (AI draft)" in xml
    assert len(list(doc.comments)) == 2
    # The unchanged text of the edited paragraph is preserved around the change.
    term = next(p for p in doc.paragraphs if "continues for two (2) years" in p.text)
    assert "The obligations in clause 3" in term.text


def test_spreadsheet_email_and_slides_are_readable():
    import email.message

    from openpyxl import Workbook
    from pptx import Presentation

    from nexus.ingest.extract import EML, PPTX, XLSX

    wb = Workbook()
    wb.active.title = "Contracts"
    wb.active.append(["Counterparty", "Change of control consent"])
    wb.active.append(["Globex Ltd", "Required (s.14.3)"])
    buf = io.BytesIO()
    wb.save(buf)
    sheet = extract(buf.getvalue(), XLSX)
    assert sheet[0].number == 1 and "Row 2: Globex Ltd | Required (s.14.3)" in sheet[0].text

    msg = email.message.EmailMessage()
    msg["From"], msg["Subject"] = "cfo@target.example", "Consent status"
    msg.set_content("We have not yet obtained Globex's consent.")
    mail = extract(msg.as_bytes(), EML)[0].text
    assert "Subject: Consent status" in mail and "not yet obtained" in mail

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Deal timeline"
    slide.placeholders[1].text = "Signing expected 30 June"
    buf = io.BytesIO()
    deck.save(buf)
    slides = extract(buf.getvalue(), PPTX)
    assert slides[0].number == 1 and "Signing expected 30 June" in slides[0].text
