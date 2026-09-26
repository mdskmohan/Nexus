"""Pull text out of uploaded files, page by page.

Supported: PDF, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), email (.eml),
plain text, Markdown and CSV — the formats that make up a real matter file
or data room.

"Pages" are real pages for PDFs, sheets for spreadsheets and slides for
decks. Word, email and text files have no pages, so they are page 1 and
passages are located by paragraph instead.
"""

import email
import email.policy
import io
import re
from dataclasses import dataclass
from html import unescape

from docx import Document as DocxDocument
from pypdf import PdfReader

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
EML = "message/rfc822"
TEXT = "text/plain"
MARKDOWN = "text/markdown"
CSV = "text/csv"

SUPPORTED = {PDF: ".pdf", DOCX: ".docx", XLSX: ".xlsx", PPTX: ".pptx", EML: ".eml",
             TEXT: ".txt", MARKDOWN: ".md", CSV: ".csv"}
PAGED = {PDF: "page", XLSX: "sheet", PPTX: "slide"}

MAX_SHEET_ROWS = 5000


class UnsupportedFile(ValueError):
    pass


class UnreadableFile(ValueError):
    pass


@dataclass
class Page:
    number: int
    text: str


def content_type_for(filename: str, declared: str | None) -> str:
    lower = filename.lower()
    for content_type, suffix in SUPPORTED.items():
        if lower.endswith(suffix):
            return content_type
    if declared in SUPPORTED:
        return declared
    raise UnsupportedFile(
        "This file type is not supported. Upload PDF, Word, Excel, PowerPoint, email (.eml) or text files."
    )


def extract(data: bytes, content_type: str) -> list[Page]:
    readers = {PDF: _pdf, DOCX: _docx, XLSX: _xlsx, PPTX: _pptx, EML: _eml}
    if content_type in readers:
        return readers[content_type](data)
    if content_type in (TEXT, MARKDOWN, CSV):
        return [Page(1, data.decode("utf-8", errors="replace"))]
    raise UnsupportedFile(content_type)


def _pdf(data: bytes) -> list[Page]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UnreadableFile("This PDF is password-protected. Remove the password and upload it again.")
        pages = [Page(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    except UnreadableFile:
        raise
    except Exception as exc:  # pypdf raises a wide range of parse errors
        raise UnreadableFile(f"This PDF could not be read ({exc.__class__.__name__}).") from exc
    if not any(p.text.strip() for p in pages):
        raise UnreadableFile(
            "This PDF has no text layer (it is probably a scan). Run it through OCR and upload it again."
        )
    return pages


def _docx(data: bytes) -> list[Page]:
    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise UnreadableFile("This Word file could not be opened.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for table_row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in table_row.cells))
    return [Page(1, "\n\n".join(parts))]


def _xlsx(data: bytes) -> list[Page]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UnreadableFile("This Excel file could not be opened.") from exc
    pages = []
    for number, sheet in enumerate(workbook.worksheets, 1):
        lines = [f"Sheet: {sheet.title}"]
        for i, values in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_SHEET_ROWS:
                lines.append(f"(sheet truncated after {MAX_SHEET_ROWS} rows)")
                break
            cells = ["" if v is None else str(v).strip() for v in values]
            if any(cells):
                # One row per paragraph, so each row can be found and cited.
                lines.append(f"Row {i + 1}: " + " | ".join(cells).rstrip(" |"))
        pages.append(Page(number, "\n\n".join(lines)))
    workbook.close()
    return pages


def _pptx(data: bytes) -> list[Page]:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise UnreadableFile("This PowerPoint file could not be opened.") from exc
    pages = []
    for number, slide in enumerate(deck.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for table_row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in table_row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text.strip())
        pages.append(Page(number, "\n\n".join(parts)))
    return pages


def _html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", "", html)
    html = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</tr>", "\n", html)
    return unescape(re.sub(r"<[^>]+>", "", html))


def _eml(data: bytes) -> list[Page]:
    try:
        message = email.message_from_bytes(data, policy=email.policy.default)
    except Exception as exc:
        raise UnreadableFile("This email file could not be read.") from exc
    header = [f"{h}: {message[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if message[h]]
    body = message.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        text = body.get_content()
        if body.get_content_type() == "text/html":
            text = _html_to_text(text)
    attachments = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
    parts = ["\n".join(header), text.strip()]
    if attachments:
        parts.append("Attachments: " + ", ".join(attachments))
    return [Page(1, "\n\n".join(p for p in parts if p))]
