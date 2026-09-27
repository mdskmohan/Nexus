"""Original, fictional Indian-law contracts with known issues.

Used for demos and for the India contract-review benchmark
(nexus/bench/india_review.py). Each case lists, for every position of its
playbook, what a careful reviewer should report:

    "issue"  the contract falls short: expect deviates, missing or unclear
    "meets"  the contract meets the standard: raising it is a false alarm
    "either" depends on facts outside the document (e.g. MSME status): any answer

The answer keys were written together with the contracts and follow the
starter playbooks. They have not yet been validated by a practising lawyer.
"""

import io

from docx import Document
from docx.shared import Pt

Clauses = list[tuple[str, list[str]]]


def contract_docx(clauses: Clauses) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for heading, paragraphs in clauses:
        if heading:
            run = doc.add_paragraph().add_run(heading)
            run.bold = True
        for text in paragraphs:
            doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


SERVICES: Clauses = [
    ("MASTER SERVICES AGREEMENT", []),
    ("", ["This Master Services Agreement (the “Agreement”) is entered into at Pune on 12 January 2026 "
          "between Kaveri Logistics Private Limited, a company incorporated under the Companies Act, 2013 with its "
          "registered office at 22 Baner Road, Pune 411045 (the “Customer”), and Sahyadri Software "
          "Solutions LLP, a limited liability partnership with its office at 9 FC Road, Pune 411004 (the "
          "“Service Provider”)."]),
    ("1. Services", [
        "1.1 The Service Provider shall develop and maintain a fleet-tracking application for the Customer as "
        "described in Schedule 1 (the “Services”).",
        "1.2 Deliverables shall be deemed accepted by the Customer upon delivery."]),
    ("2. Fees and Payment", [
        "2.1 The Customer shall pay a monthly fee of INR 4,50,000 (Rupees Four Lakh Fifty Thousand only).",
        "2.2 The Customer shall pay each invoice within 90 days of receipt.",
        "2.3 The Service Provider may revise the fees at any time by giving 15 days' notice to the Customer."]),
    ("3. Intellectual Property", [
        "3.1 The Service Provider hereby assigns to the Customer all copyright in the source code of the "
        "application developed under this Agreement."]),
    ("4. Confidentiality", [
        "4.1 Each party shall keep confidential all information received from the other party that is marked "
        "confidential and shall use it only for the purposes of this Agreement. This obligation shall survive for "
        "three years after termination."]),
    ("5. Liability and Indemnity", [
        "5.1 The Service Provider's total liability under this Agreement shall not exceed the fees paid in the "
        "three months preceding the claim.",
        "5.2 The Customer shall indemnify and hold harmless the Service Provider against any and all losses "
        "arising out of or in connection with this Agreement."]),
    ("6. Delay", [
        "6.1 If any milestone in Schedule 1 is delayed for reasons attributable to the Service Provider, the "
        "Service Provider shall pay the Customer INR 10,000 for each day of delay, up to a maximum of 10% of the "
        "fees for that milestone."]),
    ("7. Term and Termination", [
        "7.1 This Agreement commences on the date above and continues for two years.",
        "7.2 The Service Provider may terminate this Agreement at any time by giving 30 days' notice. The "
        "Customer may terminate this Agreement only for material breach by the Service Provider that is not "
        "remedied within 60 days of notice."]),
    ("8. Personal Data", [
        "8.1 The Service Provider will process location data of the Customer's drivers in providing the Services."]),
    ("9. Dispute Resolution", [
        "9.1 Any dispute arising out of this Agreement shall be referred to a sole arbitrator appointed by the "
        "Service Provider. The venue of arbitration shall be Mumbai.",
        "9.2 This Agreement is governed by the laws of India, and the courts at Delhi shall have exclusive "
        "jurisdiction."]),
    ("10. General", [
        "10.1 Neither party is liable for delay caused by events beyond its reasonable control, including flood, "
        "epidemic, war or government action, provided it notifies the other party promptly; if such an event "
        "continues for more than 60 days, either party may terminate this Agreement.",
        "10.2 This Agreement may be executed electronically."]),
]

NDA: Clauses = [
    ("MUTUAL NON-DISCLOSURE AGREEMENT", []),
    ("", ["This Mutual Non-Disclosure Agreement is made at Bengaluru on 5 February 2026 between Nilgiri Biotech "
          "Private Limited, 41 Outer Ring Road, Bengaluru 560103 (“Nilgiri”) and Coromandel Pharma "
          "Limited, 7 Anna Salai, Chennai 600002 (“Coromandel”), each a “Party”."]),
    ("1. Purpose", ["The Parties will exchange information solely to evaluate a joint development of a vaccine "
                    "cold-chain monitor (the “Purpose”)."]),
    ("2. Confidential Information", [
        "2.1 “Confidential Information” means all information disclosed by a Party to the other Party in "
        "any form, whether written, oral, visual or electronic, that is identified as confidential at the time "
        "of disclosure or within thirty (30) days after disclosure.",
        "2.2 Confidential Information does not include information that (a) is or becomes publicly available "
        "other than through breach of this Agreement, (b) was lawfully known to the recipient before disclosure, "
        "(c) is independently developed by the recipient without use of the Confidential Information, or (d) is "
        "lawfully received from a third party without restriction."]),
    ("3. Obligations", [
        "3.1 The recipient shall use Confidential Information only for the Purpose and may disclose it only to "
        "its directors, employees and professional advisers who need to know it for the Purpose and are bound by "
        "confidentiality obligations no less protective than this Agreement. The recipient is responsible for "
        "any breach by them."]),
    ("4. Term", ["This Agreement continues for one (1) year from its date, and all obligations under it end when "
                 "it expires or is terminated."]),
    ("5. Non-Competition", ["For a period of three (3) years after this Agreement ends, neither Party shall "
                            "develop, manufacture or sell any cold-chain monitoring product in India."]),
    ("6. Remedies", ["If a Party breaches this Agreement, it shall pay the other Party INR 50,00,000 (Rupees Fifty "
                     "Lakh only) as liquidated damages for each breach, without prejudice to any other remedy."]),
    ("7. Return of Information", ["On written request, the recipient shall promptly return or destroy all "
                                  "Confidential Information, except copies retained in routine backups or as "
                                  "required by law, which shall remain confidential."]),
    ("8. Governing Law and Disputes", [
        "8.1 This Agreement is governed by the laws of India.",
        "8.2 Any dispute shall be finally resolved by arbitration under the Arbitration and Conciliation Act, "
        "1996 by three arbitrators, one appointed by each Party and the third by the two so appointed. The seat "
        "of arbitration shall be Bengaluru and the language English."]),
    ("9. Stamp Duty", ["Each Party shall bear half of the stamp duty payable on this Agreement, which shall be "
                       "stamped in Karnataka before execution."]),
]

EMPLOYMENT: Clauses = [
    ("EMPLOYMENT AGREEMENT", []),
    ("", ["This Employment Agreement is made at Hyderabad on 2 March 2026 between Deccan Fintech Private Limited, "
          "Plot 12, HITEC City, Hyderabad 500081 (the “Company”) and Ms. Ananya Rao (the "
          "“Employee”)."]),
    ("1. Appointment", ["The Employee is appointed as Senior Data Engineer at the Company's Hyderabad office on an "
                        "annual cost to company of INR 32,00,000, as detailed in Annexure A. The Company may "
                        "transfer the Employee to any office or group company at any time without notice."]),
    ("2. Probation and Notice", [
        "2.1 The Employee will be on probation for six months.",
        "2.2 After confirmation, the Company may terminate employment by giving 15 days' notice. The Employee "
        "may resign by giving 90 days' notice.",
        "2.3 If the Employee leaves within two years of joining, the Employee shall pay the Company INR 5,00,000 "
        "as training costs."]),
    ("3. Restrictions", ["For two (2) years after employment ends, the Employee shall not join or work for any "
                         "company engaged in digital payments or lending anywhere in India."]),
    ("4. Confidentiality and Intellectual Property", [
        "4.1 The Employee shall keep the Company's confidential information confidential during and after "
        "employment.",
        "4.2 All software, code and other works created by the Employee in the course of employment belong to "
        "the Company, and the Employee hereby assigns to the Company all rights in them."]),
    ("5. Benefits", ["The Employee is entitled to provident fund, gratuity and leave as per applicable law and "
                     "the Company's policies, including its policy on prevention of sexual harassment."]),
    ("6. Data", ["The Company will process the Employee's personal data, including bank and health details, for "
                 "employment purposes."]),
    ("7. Disputes", ["Any dispute shall be decided by a sole arbitrator appointed by the Company."]),
]

CASES = [
    {
        "name": "Pune services agreement",
        "playbook": "services-india",
        "client_role": "We act for the Customer (Kaveri Logistics)",
        "clauses": SERVICES,
        "expected": {
            "scope": "issue",              # 1.2 deemed acceptance on delivery
            "payment_tax": "issue",        # silent on GST/TDS; 2.3 unilateral fee revision
            "msme": "either",              # depends on whether the Service Provider is a micro/small enterprise
            "liability_cap": "issue",      # 5.1 caps only the Service Provider; nothing protects the Customer
            "indemnities": "issue",        # 5.2 broad, uncapped indemnity from our client
            "liquidated_damages": "meets", # 6.1 capped LDs payable by the Service Provider
            "termination": "issue",        # 7.2 one-sided termination for convenience
            "ip": "issue",                 # 3.1 assignment silent on duration and territory (Copyright Act s.19)
            "confidentiality": "issue",    # 4.1 marked information only
            "force_majeure": "meets",      # 10.1
            "disputes": "issue",           # 9.1 unilateral appointment; venue without seat; 9.2 Delhi courts
            "stamp_execution": "issue",    # silent on stamp duty
            "personal_data": "issue",      # 8.1 drivers' location data; no DPDP terms
        },
    },
    {
        "name": "Bengaluru mutual NDA",
        "playbook": "nda-india",
        "client_role": "Mutual: our client (Nilgiri Biotech) both shares and receives information",
        "clauses": NDA,
        "expected": {
            "definition": "meets",
            "exclusions": "meets",
            "compelled_disclosure": "issue",   # no carve-out for legally compelled disclosure
            "purpose_recipients": "meets",
            "term": "issue",                   # obligations end when the agreement ends
            "restraint_of_trade": "issue",     # 3-year non-compete (s.27)
            "remedies": "issue",               # INR 50 lakh per breach liquidated damages (s.74)
            "return": "meets",
            "disputes": "meets",               # seat, three arbitrators, neutral appointment, language
            "governing_law": "meets",
            "stamp_execution": "meets",        # duty shared, stamped in Karnataka
            "personal_data": "either",         # unclear whether personal data will be exchanged
        },
    },
    {
        "name": "Hyderabad employment agreement",
        "playbook": "employment-india",
        "client_role": "We act for the Employee (Ms. Ananya Rao)",
        "clauses": EMPLOYMENT,
        "expected": {
            "role_pay": "issue",                   # transfer anywhere, any time, without notice
            "probation_notice": "issue",           # 15 vs 90 days; INR 5 lakh training bond
            "post_employment_restraint": "issue",  # 2-year post-employment non-compete (s.27)
            "confidentiality_ip": "meets",
            "statutory_benefits": "meets",
            "disputes_law": "issue",               # sole arbitrator appointed by the Company
            "personal_data": "issue",              # health and bank data; no DPDP terms
        },
    },
]
