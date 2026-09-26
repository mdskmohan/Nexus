"""Word outputs for a contract review.

1. A review memo (.docx) for every review.
2. For Word originals, a copy of the contract with each suggested change
   applied as a real tracked change (deletion + insertion, authored
   "Nexus (AI draft)") and the reasoning attached as a margin comment, so the
   lawyer accepts or rejects each change in Word as they would a colleague's.

Limitation: a changed paragraph is rebuilt with the formatting of its first
run, so mixed formatting inside that one paragraph (e.g. a bold defined term)
is simplified. Paragraphs without changes are untouched.
"""

import copy
import io
import re
from datetime import UTC, datetime

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

AUTHOR = "Nexus (AI draft)"
STATUS_LABEL = {"meets": "Meets standard", "deviates": "Deviates", "missing": "Missing",
                "unclear": "Lawyer to decide"}

_DQ, _SQ, _DASH = "[\"“”„]", "['‘’]", "[-‐-—−]"
_EQUIV = {'"': _DQ, "“": _DQ, "”": _DQ, "„": _DQ,
          "'": _SQ, "‘": _SQ, "’": _SQ,
          "-": _DASH, "–": _DASH, "—": _DASH}
_ELLIPSIS = re.compile(r"\s*(?:\.\s?\.\s?\.|…)\s*")


def _fragment_pattern(fragment: str) -> str:
    parts = []
    for word in fragment.split():
        parts.append("".join(_EQUIV.get(ch, re.escape(ch)) for ch in word))
    return r"\s+".join(parts)


def locate(quote: str, text: str) -> tuple[int, int] | None:
    """Find a (possibly ellipsis-split) quote in text; returns the covering span."""
    fragments = [f for f in _ELLIPSIS.split(quote) if f.strip()]
    if not fragments:
        return None
    start, pos = None, 0
    for fragment in fragments:
        match = re.compile(_fragment_pattern(fragment), re.IGNORECASE).search(text, pos)
        if not match:
            return None
        start = match.start() if start is None else start
        pos = match.end()
    return start, pos


class _Ids:
    def __init__(self) -> None:
        self.n = 1000

    def next(self) -> str:
        self.n += 1
        return str(self.n)


def _run(text: str, rpr, deleted: bool = False):
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:delText" if deleted else "w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _tracked(tag: str, ids: _Ids, stamp: str):
    el = OxmlElement(tag)
    el.set(qn("w:id"), ids.next())
    el.set(qn("w:author"), AUTHOR)
    el.set(qn("w:date"), stamp)
    return el


def _replace_in_paragraph(doc, paragraph: Paragraph, span: tuple[int, int], new_text: str,
                          comment: str, ids: _Ids, stamp: str) -> None:
    text = paragraph.text
    first = paragraph.runs[0]._r if paragraph.runs else None
    rpr = first.find(qn("w:rPr")) if first is not None else None
    for r in list(paragraph._p.findall(qn("w:r"))):
        paragraph._p.remove(r)
    p = paragraph._p
    before, old, after = text[: span[0]], text[span[0] : span[1]], text[span[1] :]
    if before:
        p.append(_run(before, rpr))
    deletion = _tracked("w:del", ids, stamp)
    deletion.append(_run(old, rpr, deleted=True))
    p.append(deletion)
    insertion = _tracked("w:ins", ids, stamp)
    new_r = _run(new_text, rpr)
    insertion.append(new_r)
    p.append(insertion)
    if after:
        p.append(_run(after, rpr))
    doc.add_comment(Run(new_r, paragraph), text=comment, author=AUTHOR, initials="NX")


def _append_insertion(doc, heading: str, new_text: str, comment: str, ids: _Ids, stamp: str) -> None:
    paragraph = doc.add_paragraph()
    insertion = _tracked("w:ins", ids, stamp)
    new_r = _run(new_text, None)
    insertion.append(new_r)
    paragraph._p.append(insertion)
    doc.add_comment(Run(new_r, paragraph), text=f"Proposed addition — {heading}. {comment}",
                    author=AUTHOR, initials="NX")


def tracked_changes(original: bytes, findings: list[dict]) -> tuple[bytes, dict]:
    """Apply suggested language as tracked changes. Returns the file and counts."""
    doc = Document(io.BytesIO(original))
    ids, stamp = _Ids(), datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    applied, not_placed = 0, []
    paragraphs = list(doc.paragraphs)
    for f in findings:
        suggestion = f.get("suggested_language", "")
        if not suggestion:
            continue
        note = f"{f['title']}: {f['explanation']}"
        if f["status"] == "deviates" and f.get("citation"):
            for paragraph in paragraphs:
                span = locate(f["citation"]["quote"], paragraph.text)
                if span:
                    _replace_in_paragraph(doc, paragraph, span, suggestion, note, ids, stamp)
                    applied += 1
                    break
            else:
                not_placed.append(f["title"])
        elif f["status"] == "missing":
            _append_insertion(doc, f["title"], suggestion, f["explanation"], ids, stamp)
            applied += 1
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), {"applied": applied, "not_placed": not_placed}


def _heading(doc, text: str, size: int = 13) -> None:
    run = doc.add_paragraph().add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def review_memo(result: dict, matter: str, document: str, reviewer_note: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph().add_run("Contract review memo")
    title.bold = True
    title.font.size = Pt(18)
    doc.add_paragraph(f"Matter: {matter}\nDocument: {document}\n"
                      f"Playbook: {result['playbook']['name']}\nOur client's role: {result['client_role']}")

    banner = doc.add_paragraph().add_run(
        "DRAFT prepared with AI assistance. Not legal advice until reviewed and approved by a "
        f"lawyer. {reviewer_note}"
    )
    banner.bold = True
    banner.font.color.rgb = RGBColor(0xB4, 0x3C, 0x1E)
    if result["playbook"].get("is_starter") and not result["playbook"].get("validated_by"):
        doc.add_paragraph(
            "This review used a Nexus starter playbook that has not yet been approved by a lawyer "
            "at the firm. Its positions are general commercial practice, not advice on any "
            "jurisdiction's law."
        )

    _heading(doc, "Key points")
    for line in [x.strip() for x in result["summary"].splitlines() if x.strip()]:
        doc.add_paragraph(line.lstrip("-•* "), style="List Bullet")

    _heading(doc, "Findings")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, ["Position", "Result", "Risk", "Contract language"], strict=True):
        cell.text = label
    for f in result["findings"]:
        cells = table.add_row().cells
        cells[0].text = f["title"]
        cells[1].text = STATUS_LABEL[f["status"]]
        cells[2].text = f["risk"].capitalize()
        c = f.get("citation")
        cells[3].text = f"“{c['quote']}” (page {c['page']})" if c else "—"

    _heading(doc, "Detail and suggested language")
    for f in result["findings"]:
        if f["status"] == "meets":
            continue
        p = doc.add_paragraph()
        p.add_run(f"{f['title']} — {STATUS_LABEL[f['status']]} ({f['risk']} risk)").bold = True
        doc.add_paragraph(f["explanation"])
        if f["suggested_language"]:
            s = doc.add_paragraph()
            s.add_run("Suggested language: ").italic = True
            s.add_run(f["suggested_language"])

    _heading(doc, "Draft note to client")
    doc.add_paragraph(result["client_note"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.LINE)
    doc.add_paragraph("Reviewed and approved by: ______________________   Date: ____________")

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
